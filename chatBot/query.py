import os
import uuid
from typing import List, Dict, Any, Optional
import numpy as np
from qdrant_client import QdrantClient
from qdrant_client.http import models
from qdrant_client.http.models import Distance, VectorParams
from sentence_transformers import SentenceTransformer
import logging
from pathlib import Path
import mimetypes
import docx
from PyPDF2 import PdfReader
import json

# Configuration du logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class QdrantEmbeddedVectorStore:
    """
    Classe pour gérer un vectorstore avec Qdrant en mode embedded (sans serveur)
    """
    
    def __init__(
        self,
        collection_name: str = "documents",
        db_path: str = "./qdrant_db",
        embedding_model: str = "all-MiniLM-L6-v2"
    ):
        """
        Initialise Qdrant en mode embedded
        
        Args:
            collection_name: Nom de la collection Qdrant
            db_path: Chemin local pour stocker la base de données
            embedding_model: Nom du modèle Sentence Transformers
        """
        self.collection_name = collection_name
        self.db_path = db_path
        
        # Créer le dossier de la base de données si nécessaire
        os.makedirs(db_path, exist_ok=True)
        
        # Initialisation du client Qdrant en mode embedded
        self.client = QdrantClient(path=db_path)
        logger.info(f"Qdrant initialisé en mode embedded dans: {db_path}")
        
        # Initialisation du modèle d'embedding
        self.embedding_model = SentenceTransformer(embedding_model)
        self.vector_size = self.embedding_model.get_sentence_embedding_dimension()
        
        # Créer la collection si elle n'existe pas
        self._create_collection_if_not_exists()
    
    def _create_collection_if_not_exists(self):
        """Crée la collection Qdrant si elle n'existe pas"""
        try:
            collections = self.client.get_collections().collections
            collection_names = [col.name for col in collections]
            
            if self.collection_name not in collection_names:
                self.client.create_collection(
                    collection_name=self.collection_name,
                    vectors_config=VectorParams(
                        size=self.vector_size,
                        distance=Distance.COSINE
                    )
                )
                logger.info(f"Collection '{self.collection_name}' créée avec succès")
            else:
                logger.info(f"Collection '{self.collection_name}' existe déjà")
                
        except Exception as e:
            logger.error(f"Erreur lors de la création de la collection: {e}")
            raise
    
    def _extract_text_from_file(self, file_path: str) -> tuple[str, Dict[str, Any]]:
        """
        Extrait le texte d'un fichier selon son type
        
        Args:
            file_path: Chemin vers le fichier
            
        Returns:
            Tuple (contenu_texte, métadonnées_fichier)
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"Le fichier {file_path} n'existe pas")
        
        # Métadonnées de base du fichier
        file_metadata = {
            "filename": file_path.name,
            "file_path": str(file_path.absolute()),
            "file_size": file_path.stat().st_size,
            "file_extension": file_path.suffix.lower(),
            "mime_type": mimetypes.guess_type(str(file_path))[0]
        }
        
        try:
            # Extraction selon le type de fichier
            if file_path.suffix.lower() == '.txt':
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                    
            elif file_path.suffix.lower() == '.pdf':
                content = ""
                with open(file_path, 'rb') as f:
                    pdf_reader = PdfReader(f)
                    for page in pdf_reader.pages:
                        content += page.extract_text() + "\n"
                file_metadata["num_pages"] = len(pdf_reader.pages)
                
            elif file_path.suffix.lower() in ['.docx', '.doc']:
                doc = docx.Document(file_path)
                content = ""
                for paragraph in doc.paragraphs:
                    content += paragraph.text + "\n"
                file_metadata["num_paragraphs"] = len(doc.paragraphs)
                
            elif file_path.suffix.lower() == '.json':
                with open(file_path, 'r', encoding='utf-8') as f:
                    json_data = json.load(f)
                    content = json.dumps(json_data, indent=2, ensure_ascii=False)
                    
            elif file_path.suffix.lower() in ['.md', '.markdown']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                    
            elif file_path.suffix.lower() in ['.py', '.js', '.html', '.css', '.cpp', '.java', '.c']:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                file_metadata["file_type"] = "code"
                
            else:
                # Essayer de lire comme fichier texte
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        content = f.read()
                except UnicodeDecodeError:
                    try:
                        with open(file_path, 'r', encoding='latin-1') as f:
                            content = f.read()
                    except:
                        raise ValueError(f"Impossible de lire le fichier {file_path}. Type non supporté.")
            
            # Nettoyer le contenu
            content = content.strip()
            if not content:
                raise ValueError(f"Le fichier {file_path} est vide ou ne contient pas de texte extractible")
            
            return content, file_metadata
            
        except Exception as e:
            logger.error(f"Erreur lors de l'extraction du fichier {file_path}: {e}")
            raise

    def add_documents(
        self,
        file_paths: List[str],
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
        additional_metadata: Optional[List[Dict[str, Any]]] = None,
        ids: Optional[List[str]] = None
    ) -> List[str]:
        """
        Ajoute des documents au vectorstore à partir de chemins de fichiers
        
        Args:
            file_paths: Liste des chemins vers les fichiers
            chunk_size: Taille des chunks de texte (caractères)
            chunk_overlap: Chevauchement entre chunks (caractères)
            additional_metadata: Métadonnées supplémentaires par fichier (optionnel)
            ids: Liste des IDs personnalisés (optionnel)
            
        Returns:
            Liste des IDs des documents ajoutés
        """
        if not file_paths:
            raise ValueError("La liste des chemins de fichiers ne peut pas être vide")
        
        all_documents = []
        all_metadatas = []
        all_ids = []
        
        # Traiter chaque fichier
        for i, file_path in enumerate(file_paths):
            try:
                logger.info(f"Traitement du fichier: {file_path}")
                
                # Extraire le texte et les métadonnées du fichier
                content, file_metadata = self._extract_text_from_file(file_path)
                
                # Ajouter les métadonnées supplémentaires si fournies
                if additional_metadata and i < len(additional_metadata):
                    file_metadata.update(additional_metadata[i])
                
                # Découper le contenu en chunks si nécessaire
                chunks = self._split_text(content, chunk_size, chunk_overlap)
                
                # Créer les documents et métadonnées pour chaque chunk
                for chunk_idx, chunk in enumerate(chunks):
                    chunk_metadata = file_metadata.copy()
                    chunk_metadata.update({
                        "chunk_index": chunk_idx,
                        "total_chunks": len(chunks),
                        "chunk_size": len(chunk)
                    })
                    
                    # Générer un ID unique pour ce chunk
                    if ids and i < len(ids):
                        chunk_id = f"{ids[i]}_chunk_{chunk_idx}" if len(chunks) > 1 else ids[i]
                    else:
                        base_id = str(uuid.uuid4())
                        chunk_id = f"{base_id}_chunk_{chunk_idx}" if len(chunks) > 1 else base_id
                    
                    all_documents.append(chunk)
                    all_metadatas.append(chunk_metadata)
                    all_ids.append(chunk_id)
                
                logger.info(f"Fichier {file_path} traité: {len(chunks)} chunks créés")
                
            except Exception as e:
                logger.error(f"Erreur lors du traitement du fichier {file_path}: {e}")
                raise
        
        try:
            # Générer les embeddings
            logger.info(f"Génération des embeddings pour {len(all_documents)} chunks...")
            embeddings = self.embedding_model.encode(all_documents, show_progress_bar=True)
            
            # Préparer les points pour Qdrant
            points = []
            for doc_id, doc, metadata, embedding in zip(all_ids, all_documents, all_metadatas, embeddings):
                # Ajouter le texte aux métadonnées
                full_metadata = {**metadata, "text": doc}
                
                point = models.PointStruct(
                    id=doc_id,
                    vector=embedding.tolist(),
                    payload=full_metadata
                )
                points.append(point)
            
            # Uploader les points vers Qdrant (par batches si nécessaire)
            batch_size = 100
            for i in range(0, len(points), batch_size):
                batch = points[i:i + batch_size]
                self.client.upsert(
                    collection_name=self.collection_name,
                    points=batch
                )
                logger.info(f"Batch {i//batch_size + 1}/{(len(points)-1)//batch_size + 1} uploadé")
            
            logger.info(f"{len(all_documents)} chunks ajoutés avec succès depuis {len(file_paths)} fichiers")
            return all_ids
            
        except Exception as e:
            logger.error(f"Erreur lors de l'ajout des documents: {e}")
            raise

    def _split_text(self, text: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[str]:
        """
        Découpe un texte en chunks avec chevauchement
        
        Args:
            text: Texte à découper
            chunk_size: Taille maximale d'un chunk
            chunk_overlap: Chevauchement entre chunks
            
        Returns:
            Liste des chunks
        """
        if len(text) <= chunk_size:
            return [text]
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + chunk_size
            
            # Si ce n'est pas le dernier chunk, essayer de couper à un espace
            if end < len(text):
                # Chercher le dernier espace dans la zone de chevauchement
                last_space = text.rfind(' ', start + chunk_size - chunk_overlap, end)
                if last_space > start:
                    end = last_space
            
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            # Calculer le nouveau point de départ avec chevauchement
            start = end - chunk_overlap
            if start >= len(text):
                break
        
        return chunks
    
    def similarity_search(
        self,
        query: str,
        k: int = 5,
        filter_conditions: Optional[Dict[str, Any]] = None
    ) -> List[Dict[str, Any]]:
        """
        Recherche par similarité dans le vectorstore
        
        Args:
            query: Texte de la requête
            k: Nombre de résultats à retourner
            filter_conditions: Conditions de filtrage (optionnel)
            
        Returns:
            Liste des résultats avec scores et métadonnées
        """
        try:
            # Générer l'embedding de la requête
            query_embedding = self.embedding_model.encode([query])[0]
            
            # Préparer le filtre si fourni
            query_filter = None
            if filter_conditions:
                query_filter = models.Filter(
                    must=[
                        models.FieldCondition(
                            key=key,
                            match=models.MatchValue(value=value)
                        )
                        for key, value in filter_conditions.items()
                    ]
                )
            
            # Effectuer la recherche
            search_results = self.client.search(
                collection_name=self.collection_name,
                query_vector=query_embedding.tolist(),
                limit=k,
                query_filter=query_filter,
                with_payload=True,
                with_vectors=False
            )
            
            # Formater les résultats
            results = []
            for result in search_results:
                results.append({
                    "id": result.id,
                    "text": result.payload.get("text", ""),
                    "metadata": {k: v for k, v in result.payload.items() if k != "text"},
                    "score": result.score
                })
            
            return results
            
        except Exception as e:
            logger.error(f"Erreur lors de la recherche: {e}")
            raise
    
    def delete_documents(self, ids: List[str]) -> bool:
        """
        Supprime des documents du vectorstore
        
        Args:
            ids: Liste des IDs à supprimer
            
        Returns:
            True si la suppression a réussi
        """
        try:
            self.client.delete(
                collection_name=self.collection_name,
                points_selector=models.PointIdsList(
                    points=ids
                )
            )
            logger.info(f"{len(ids)} documents supprimés")
            return True
            
        except Exception e:
            logger.error(f"Erreur lors de la suppression: {e}")
            return False
    
    def get_collection_info(self) -> Dict[str, Any]:
        """Retourne les informations sur la collection"""
        try:
            info = self.client.get_collection(self.collection_name)
            return {
                "name": self.collection_name,
                "vector_size": info.config.params.vectors.size,
                "vectors_count": info.vectors_count,
                "points_count": info.points_count,
                "status": info.status,
                "db_path": self.db_path
            }
        except Exception as e:
            logger.error(f"Erreur lors de la récupération des infos: {e}")
            return {}
    
    def backup_database(self, backup_path: str):
        """Crée une sauvegarde de la base de données"""
        import shutil
        try:
            shutil.copytree(self.db_path, backup_path)
            logger.info(f"Sauvegarde créée dans: {backup_path}")
            return True
        except Exception as e:
            logger.error(f"Erreur lors de la sauvegarde: {e}")
            return False

# Exemple d'utilisation
def example_usage():
    """Exemple d'utilisation du QdrantEmbeddedVectorStore"""
    
    print("🚀 Initialisation de Qdrant en mode embedded...")
    
    # Initialiser le vectorstore embedded
    vectorstore = QdrantEmbeddedVectorStore(
        collection_name="ma_collection_embedded",
        db_path="./mon_qdrant_db",  # Base de données locale
        embedding_model="all-MiniLM-L6-v2"
    )
    
    print("✅ Qdrant embedded initialisé avec succès!")
    
    # Créer des fichiers de test
    test_files = []
    
    # Créer un fichier texte de test
    with open("test_doc.txt", "w", encoding="utf-8") as f:
        f.write("Ceci est un document de test pour Qdrant embedded. "
                "Il contient des informations importantes sur l'intelligence artificielle.")
    test_files.append("test_doc.txt")
    
    # Créer un fichier JSON de test
    with open("test_config.json", "w", encoding="utf-8") as f:
        json.dump({
            "name": "Mon Application",
            "version": "1.0.0",
            "description": "Une application qui utilise Qdrant embedded pour la recherche vectorielle"
        }, f, indent=2, ensure_ascii=False)
    test_files.append("test_config.json")
    
    # Métadonnées supplémentaires
    additional_metadata = [
        {"category": "documentation", "author": "Test User"},
        {"category": "config", "environment": "development"}
    ]
    
    print("📁 Ajout des fichiers...")
    try:
        doc_ids = vectorstore.add_documents(
            file_paths=test_files,
            chunk_size=500,
            chunk_overlap=50,
            additional_metadata=additional_metadata
        )
        print(f"✅ {len(doc_ids)} chunks ajoutés avec succès!")
        
    except Exception as e:
        print(f"❌ Erreur lors de l'ajout: {e}")
        return
    
    print("\n🔍 Test de recherche...")
    query = "intelligence artificielle application"
    results = vectorstore.similarity_search(query, k=3)
    
    print(f"Résultats pour: '{query}'")
    for i, result in enumerate(results, 1):
        print(f"\n{i}. Score: {result['score']:.3f}")
        print(f"   Fichier: {result['metadata'].get('filename', 'N/A')}")
        print(f"   Chunk: {result['metadata'].get('chunk_index', 0)}")
        print(f"   Extrait: {result['text'][:100]}...")
    
    print("\n📊 Informations sur la collection:")
    info = vectorstore.get_collection_info()
    for key, value in info.items():
        print(f"   {key}: {value}")
    
    print("\n💾 Création d'une sauvegarde...")
    backup_success = vectorstore.backup_database("./backup_qdrant_db")
    if backup_success:
        print("✅ Sauvegarde créée avec succès!")
    
    # Nettoyer les fichiers de test
    for file_path in test_files:
        if os.path.exists(file_path):
            os.remove(file_path)
    
    print("\n🎉 Test terminé avec succès!")

if __name__ == "__main__":
    example_usage()
